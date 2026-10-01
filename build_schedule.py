# -*- coding: utf-8 -*-
"""schedule/일정3.xlsx 를 읽어 schedule-data.js 를 만든다.

사용법 (저장소 루트에서):
    pip install openpyxl
    python3 scripts/build_schedule.py

엑셀의 "전체 일정" 시트 A열 배경색으로 구분(실험실/대면/구글챗·개인/마감·주문/행사)을 나눈다.
"""
import json
import pathlib
import re
import sys

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / "schedule" / "일정3.xlsx"
OUT = ROOT / "schedule-data.js"
YEAR = 2026

# 행 배경색 -> 구분 코드
FILL_TO_CAT = {
    "FFFCE4D6": "lab",    # 주황: 실험실
    "FFDDEBF7": "meet",   # 파랑: 대면
    "FFF2F2F2": "solo",   # 회색: 구글챗·개인
    "FFFFE699": "due",    # 노랑: 마감·주문
    "FFC6E0B4": "event",  # 초록: 행사
}


def text(v):
    if v is None:
        return ""
    return str(v).replace("\r", "").strip()


def iso(md):
    m = re.match(r"^(\d{1,2})/(\d{1,2})$", md)
    if not m:
        raise ValueError("날짜 형식이 이상합니다: %r" % md)
    return "%d-%02d-%02d" % (YEAR, int(m.group(1)), int(m.group(2)))


def cat_of(cell):
    fill = cell.fill
    rgb = fill.fgColor.rgb if fill is not None and fill.fill_type else None
    return FILL_TO_CAT.get(rgb, "solo")


def read_schedule(ws):
    items = []
    r = 7  # 6행이 머리글
    while True:
        md = text(ws.cell(r, 1).value)
        if not re.match(r"^\d{1,2}/\d{1,2}$", md):
            break
        items.append({
            "key": md + "|" + text(ws.cell(r, 4).value),
            "date": iso(md),
            "md": md,
            "wd": text(ws.cell(r, 2).value),
            "dday": text(ws.cell(r, 3).value),
            "slot": text(ws.cell(r, 4).value),
            "form": text(ws.cell(r, 5).value),
            "move": text(ws.cell(r, 6).value),
            "place": text(ws.cell(r, 7).value),
            "task": text(ws.cell(r, 8).value),
            "owner": text(ws.cell(r, 9).value),
            "cat": cat_of(ws.cell(r, 1)),
        })
        r += 1
    return {
        "title": text(ws["A1"].value),
        "subtitle": text(ws["A2"].value),
        "rule": text(ws["A3"].value),
        "legend": text(ws["A4"].value),
        "items": items,
    }


def read_moves(ws):
    rows = []
    r = 5
    while text(ws.cell(r, 1).value).endswith("주차"):
        rows.append({
            "week": text(ws.cell(r, 1).value),
            "md": text(ws.cell(r, 2).value),
            "wd": text(ws.cell(r, 3).value),
            "slot": text(ws.cell(r, 4).value),
            "place": text(ws.cell(r, 5).value),
            "what": text(ws.cell(r, 6).value),
            "due": text(ws.cell(r, 7).value),
            "cat": cat_of(ws.cell(r, 1)),
        })
        r += 1
    return {
        "title": text(ws["A1"].value),
        "note": text(ws["A2"].value),
        "rows": rows,
        "noneTitle": text(ws["A20"].value),
        "noneText": text(ws["A21"].value),
        "footnote": text(ws["A24"].value),
    }


def read_lab(ws):
    sessions = []
    r = 5
    while text(ws.cell(r, 1).value).isdigit():
        sessions.append({
            "n": text(ws.cell(r, 1).value),
            "md": text(ws.cell(r, 2).value),
            "wd": text(ws.cell(r, 3).value),
            "slot": text(ws.cell(r, 4).value),
            "purpose": text(ws.cell(r, 5).value),
            "equip": text(ws.cell(r, 6).value),
            "kind": text(ws.cell(r, 7).value),
        })
        r += 1
    checks = []
    r = 14
    while text(ws.cell(r, 1).value):
        checks.append({
            "n": text(ws.cell(r, 1).value),
            "item": text(ws.cell(r, 2).value),
            "why": text(ws.cell(r, 5).value),
        })
        r += 1
    return {
        "title": text(ws["A1"].value),
        "note": text(ws["A2"].value),
        "sessions": sessions,
        "checksTitle": text(ws["A12"].value),
        "checks": checks,
    }


def main():
    wb = openpyxl.load_workbook(XLSX)
    sched = read_schedule(wb["전체 일정"])
    data = {
        "event": {"date": "%d-11-07" % YEAR, "name": "학술교류 포럼 발표", "place": "북일고등학교"},
        "schedule": sched,
        "moves": read_moves(wb["장소이동신청"]),
        "lab": read_lab(wb["실험실 예약"]),
    }
    n = len(sched["items"])
    if n == 0:
        sys.exit("일정 행을 하나도 읽지 못했습니다. 엑셀 형식을 확인하세요.")
    body = json.dumps(data, ensure_ascii=False, indent=1)
    OUT.write_text("// scripts/build_schedule.py 가 만든 파일입니다. 직접 고치지 마세요.\n"
                   "window.SCHEDULE = " + body + ";\n", encoding="utf-8")
    print("일정 %d칸, 장소이동신청 %d줄, 실험실 %d회차·확인 %d항목 -> %s" % (
        n, len(data["moves"]["rows"]), len(data["lab"]["sessions"]), len(data["lab"]["checks"]), OUT.name))


if __name__ == "__main__":
    main()
